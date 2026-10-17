
import io
import json
import os
import re
import uuid
from collections import OrderedDict
from pathlib import Path

import requests
from docx import Document
from flask import Flask, Response, jsonify, request, send_from_directory, stream_with_context
from pypdf import PdfReader

BASE_DIR = Path(__file__).parent.resolve()
OLLAMA_URL = os.environ.get("OLLAMA_URL", "http://localhost:11434")
DEFAULT_MODEL = os.environ.get("OFFLINEMIND_MODEL", "llama3.2")
NUM_CTX = int(os.environ.get("OFFLINEMIND_NUM_CTX", "8192"))

CONTEXT_CHARS = 12000       # max document text sent in one prompt
CHUNK_CHARS = 8000          # chunk size for map-reduce summaries
MAX_CHUNKS = 30
MAX_DOCS = 20
ALLOWED_EXT = {".pdf", ".docx", ".txt", ".md", ".csv", ".json", ".py", ".js", ".html", ".log"}

app = Flask(__name__, static_folder=None)
app.config["MAX_CONTENT_LENGTH"] = 40 * 1024 * 1024  # 40 MB
DOCS = OrderedDict()  # doc_id -> {"name": str, "text": str}


# --------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------
def extract_text(data: bytes, ext: str) -> str:
    if ext == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join((p.extract_text() or "") for p in reader.pages)
    if ext == ".docx":
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    return data.decode("utf-8", errors="ignore")


def clean_text(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def chunk_text(text: str, size: int):
    chunks, current = [], ""
    for para in text.split("\n"):
        if len(current) + len(para) + 1 > size and current:
            chunks.append(current)
            current = ""
        while len(para) > size:
            chunks.append(para[:size])
            para = para[size:]
        current += para + "\n"
    if current.strip():
        chunks.append(current)
    return chunks


def spread_sample(chunks, limit_chars):
    """Pick evenly spaced chunks so the whole document is represented."""
    total = sum(len(c) for c in chunks)
    if total <= limit_chars:
        return "\n".join(chunks)
    n = max(1, limit_chars // max(1, len(chunks[0])))
    step = len(chunks) / n
    return "\n".join(chunks[int(i * step)] for i in range(n))


def relevant_context(text: str, query: str, limit_chars: int) -> str:
    """Simple offline retrieval: keyword-overlap ranking of chunks."""
    if len(text) <= limit_chars:
        return text
    chunks = chunk_text(text, 1500)
    words = {w for w in re.findall(r"[a-zA-Z0-9]{3,}", query.lower())}
    scored = []
    for i, c in enumerate(chunks):
        cw = re.findall(r"[a-zA-Z0-9]{3,}", c.lower())
        score = sum(1 for w in cw if w in words)
        scored.append((score, i))
    scored.sort(reverse=True)
    picked, used = [], 0
    for score, i in scored:
        if used + len(chunks[i]) > limit_chars:
            continue
        picked.append(i)
        used += len(chunks[i])
    if not picked:
        picked = [0]
    return "\n...\n".join(chunks[i] for i in sorted(picked))


def ollama_stream(messages, model, temperature=0.7):
    payload = {
        "model": model,
        "messages": messages,
        "stream": True,
        "options": {"temperature": temperature, "num_ctx": NUM_CTX},
    }
    with requests.post(f"{OLLAMA_URL}/api/chat", json=payload, stream=True, timeout=(5, 900)) as r:
        if r.status_code != 200:
            try:
                msg = r.json().get("error", r.text)
            except Exception:
                msg = r.text
            raise RuntimeError(msg)
        for line in r.iter_lines():
            if not line:
                continue
            data = json.loads(line)
            if "error" in data:
                raise RuntimeError(data["error"])
            token = data.get("message", {}).get("content", "")
            if token:
                yield token
            if data.get("done"):
                break


def ollama_complete(messages, model, temperature=0.3):
    return "".join(ollama_stream(messages, model, temperature))


def line(obj):
    return json.dumps(obj) + "\n"


def stream_response(generator_fn):
    def generate():
        try:
            for item in generator_fn():
                yield line(item if isinstance(item, dict) else {"token": item})
            yield line({"done": True})
        except requests.exceptions.ConnectionError:
            yield line({"error": "Cannot reach Ollama. Start it with `ollama serve` and try again."})
        except Exception as e:  # noqa: BLE001
            yield line({"error": str(e)})

    return Response(
        stream_with_context(generate()),
        mimetype="application/x-ndjson",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


def get_doc(doc_id):
    doc = DOCS.get(doc_id)
    if not doc:
        raise RuntimeError("This document is no longer loaded (the server restarted). Please upload it again.")
    return doc


# --------------------------------------------------------------------------
# Routes
# --------------------------------------------------------------------------
@app.get("/")
def index():
    return send_from_directory(BASE_DIR, "index.html")


@app.get("/<path:filename>")
def assets(filename):
    if filename in {"style.css", "script.js"}:
        return send_from_directory(BASE_DIR, filename)
    return ("Not found", 404)


@app.get("/api/models")
def models():
    try:
        r = requests.get(f"{OLLAMA_URL}/api/tags", timeout=4)
        names = [m["name"] for m in r.json().get("models", [])]
        return jsonify({"ok": True, "models": names, "default": DEFAULT_MODEL})
    except Exception:  # noqa: BLE001
        return jsonify({"ok": False, "models": [], "default": DEFAULT_MODEL})


@app.post("/api/upload")
def upload():
    f = request.files.get("file")
    if not f or not f.filename:
        return jsonify({"error": "No file received."}), 400
    ext = Path(f.filename).suffix.lower()
    if ext not in ALLOWED_EXT:
        return jsonify({"error": f"Unsupported file type {ext or '(none)'}. Use PDF, DOCX, TXT or MD."}), 400
    try:
        text = clean_text(extract_text(f.read(), ext))
    except Exception as e:  # noqa: BLE001
        return jsonify({"error": f"Could not read the file: {e}"}), 400
    if len(text) < 20:
        return jsonify({"error": "No readable text found (scanned PDFs need OCR first)."}), 400
    doc_id = uuid.uuid4().hex[:12]
    DOCS[doc_id] = {"name": f.filename, "text": text}
    while len(DOCS) > MAX_DOCS:
        DOCS.popitem(last=False)
    return jsonify({"id": doc_id, "name": f.filename, "chars": len(text), "words": len(text.split())})


@app.delete("/api/docs/<doc_id>")
def delete_doc(doc_id):
    DOCS.pop(doc_id, None)
    return jsonify({"ok": True})


@app.post("/api/chat")
def chat():
    body = request.get_json(force=True)
    model = body.get("model") or DEFAULT_MODEL
    history = [m for m in body.get("messages", []) if m.get("role") in ("user", "assistant")]
    doc_id = body.get("doc_id")

    def gen():
        system = (
            "You are OfflineMindAI, a helpful, accurate assistant running locally on the user's computer. "
            "Answer clearly and use Markdown for formatting. Be honest when you are unsure."
        )
        if doc_id:
            doc = get_doc(doc_id)
            last_user = next((m["content"] for m in reversed(history) if m["role"] == "user"), "")
            ctx = relevant_context(doc["text"], last_user, CONTEXT_CHARS)
            system += (
                f"\n\nThe user uploaded a document named \"{doc['name']}\". Answer using it when relevant "
                f"and say so if the answer is not in it.\n\n--- DOCUMENT ---\n{ctx}\n--- END ---"
            )
        yield from ollama_stream([{"role": "system", "content": system}] + history[-20:], model)

    return stream_response(gen)


@app.post("/api/summarize")
def summarize():
    body = request.get_json(force=True)
    model = body.get("model") or DEFAULT_MODEL
    style = body.get("style", "medium")
    instructions = {
        "short": "Write a concise summary of 4-6 sentences.",
        "medium": "Write a clear summary: a short overview paragraph, then 5-8 key points as bullets.",
        "detailed": "Write a detailed structured summary with headings for each major section, key facts, "
                    "definitions and conclusions.",
    }.get(style, "Write a clear summary with key points.")

    def gen():
        doc = get_doc(body.get("doc_id"))
        text = doc["text"]
        if len(text) <= CONTEXT_CHARS:
            material = text
        else:
            chunks = chunk_text(text, CHUNK_CHARS)
            if len(chunks) > MAX_CHUNKS:
                step = len(chunks) / MAX_CHUNKS
                chunks = [chunks[int(i * step)] for i in range(MAX_CHUNKS)]
            notes = []
            for i, c in enumerate(chunks, 1):
                yield {"status": f"Reading part {i} of {len(chunks)}…"}
                notes.append(ollama_complete([
                    {"role": "system", "content": "You write compact, factual study notes."},
                    {"role": "user", "content": f"Write concise bullet notes of the key information in this text:\n\n{c}"},
                ], model))
            material = "\n\n".join(notes)[:CONTEXT_CHARS * 2]
        yield {"status": "Writing summary…"}
        yield from ollama_stream([
            {"role": "system", "content": "You are an expert summarizer. Use Markdown. Only use facts from the text."},
            {"role": "user", "content": f"{instructions}\n\nDocument: \"{doc['name']}\"\n\n{material}"},
        ], model, 0.3)

    return stream_response(gen)


@app.post("/api/exam")
def exam():
    body = request.get_json(force=True)
    model = body.get("model") or DEFAULT_MODEL
    count = max(1, min(int(body.get("count", 10)), 40))
    qtype = body.get("qtype", "mcq")
    difficulty = body.get("difficulty", "medium")
    type_rules = {
        "mcq": "multiple-choice questions, each with 4 options labelled A) B) C) D) and exactly one correct answer",
        "truefalse": "True/False questions",
        "short": "short-answer questions (1-3 sentence answers)",
        "mixed": "a mix of multiple-choice (A-D), True/False, and short-answer questions",
    }.get(qtype, "multiple-choice questions")

    def gen():
        doc = get_doc(body.get("doc_id"))
        yield {"status": "Reading the document…"}
        material = spread_sample(chunk_text(doc["text"], 2000), CONTEXT_CHARS)
        yield {"status": "Writing your exam…"}
        prompt = (
            f"Create a {difficulty}-difficulty exam of exactly {count} {type_rules}, based ONLY on the document below.\n\n"
            "Format in Markdown:\n"
            f"# Exam: {doc['name']}\n"
            "Then number the questions 1..N. Do not reveal answers next to the questions.\n"
            "After all questions, add a section '## Answer Key' listing the correct answer for every question "
            "with a one-line explanation.\n\n"
            f"--- DOCUMENT ---\n{material}\n--- END ---"
        )
        yield from ollama_stream([
            {"role": "system", "content": "You are an experienced teacher who writes fair, clear exams."},
            {"role": "user", "content": prompt},
        ], model, 0.5)

    return stream_response(gen)


if __name__ == "__main__":
    print("OfflineMindAI running at http://127.0.0.1:5000  (Ollama:", OLLAMA_URL, ")")
    app.run(host="127.0.0.1", port=5000, threaded=True, debug=False)